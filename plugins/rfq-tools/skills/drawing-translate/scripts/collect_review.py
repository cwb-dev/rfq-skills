"""Merge the <out>_review.json files of one or more translated drawings into a review spreadsheet
for a colleague: original, translation, why it needs checking, and blank columns for their verdict.

Usage: python collect_review.py <out.xlsx> <review.json | folder> [...] [--title "RFQ 1234 图纸翻译"]
A folder is searched for *_review.json. Drawings without a review file have nothing to check.
"""
import glob
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

args = sys.argv[1:]
title = "图纸中文翻译"
if "--title" in args:
    i = args.index("--title")
    title = args[i + 1]
    del args[i:i + 2]
OUT, srcs = args[0], args[1:]
files = []
for s in srcs:
    files += sorted(glob.glob(os.path.join(s, "*_review.json"))) if os.path.isdir(s) else [s]

rows = []
for f in files:
    r = json.load(open(f, encoding="utf-8-sig"))
    for it in r["items"]:
        rows.append([os.path.splitext(r["drawing"])[0], it.get("page", 1), it.get("en", ""),
                     it.get("zh") or "（未翻译）", it["review"]])

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "待审核"
thin = Side(style="thin")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
ws.merge_cells("A1:H1")
ws["A1"] = f"{title} —— 待审核清单（请核对译文，在\"审核意见\"栏填写：OK / 修改为……）"
ws["A1"].font = Font(name="微软雅黑", size=12, bold=True)
head = ["序号", "图纸", "页", "原文", "译文", "待确认原因", "审核意见", "审核人"]
for c, (h, w) in enumerate(zip(head, [6, 30, 5, 40, 40, 40, 30, 10]), 1):
    cell = ws.cell(2, c, h)
    cell.font = Font(name="微软雅黑", size=10, bold=True)
    cell.fill = PatternFill("solid", fgColor="FFFFFF00" if h in ("审核意见", "审核人") else "FFD9D9D9")
    cell.border = B
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions[chr(64 + c)].width = w
for i, row in enumerate(rows, 1):
    for c, v in enumerate([i] + row + ["", ""], 1):
        cell = ws.cell(i + 2, c, v)
        cell.font = Font(name="微软雅黑", size=10)
        cell.border = B
        cell.alignment = Alignment(wrap_text=True, vertical="center",
                                   horizontal="center" if c in (1, 3) else "left")
    n = max(len(str(v)) // 22 + str(v).count("\n") + 1 for v in row)
    ws.row_dimensions[i + 2].height = max(30, min(n, 12) * 15)
ws.freeze_panes = "A3"
ws.page_setup.orientation = "landscape"
wb.save(OUT)
print(f"{len(rows)} item(s) from {len(files)} drawing(s) -> {OUT}")
