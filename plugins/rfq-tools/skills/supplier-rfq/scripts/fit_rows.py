"""Final pass: size the 询价单 sheet's item rows to fit the wrapped text."""
import sys

import openpyxl

W = {3: 30, 5: 26, 6: 36, 11: 42}


def lines(text, width):
    n = 0
    for part in str(text).split("\n"):
        w = sum(2 if ord(ch) > 0x2E80 else 1 for ch in part)
        n += max(1, -(-w // max(int(width * 1.0), 1)))
    return n


for f in sys.argv[1:]:
    wb = openpyxl.load_workbook(f)
    ws = wb.worksheets[0]
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 1).value != "=ROW()-2":
            continue
        n = max(lines(ws.cell(r, c).value or "", w) for c, w in W.items())
        ws.row_dimensions[r].height = max(30, n * 15 + 6)
    wb.save(f)
    print("fitted", f)
