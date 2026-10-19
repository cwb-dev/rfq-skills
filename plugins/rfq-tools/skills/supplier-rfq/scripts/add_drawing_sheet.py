"""Add a drawing sheet (named by drawing no.) to a 询价单 xlsx: full drawing images + Chinese notes table.

Usage: python add_drawing_sheet.py <询价单.xlsx> <drawing.pdf> <notes.json> <img_dir>

notes.json: see ../example_notes.json
Re-running replaces the sheet, so edit the JSON and run again.
"""
import json
import sys

import openpyxl
import pymupdf
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.hyperlink import Hyperlink
from PIL import Image as PILImage

XLSX, PDF, NOTES, IMGDIR = sys.argv[1:5]
spec = json.load(open(NOTES, encoding="utf-8"))
DRW_NO = spec["drawing_no"]
SHEET = DRW_NO[:31].translate(str.maketrans({c: "-" for c in "\\/?*[]:"}))  # Excel sheet-name rules
FONT = "微软雅黑"
thin = Side(style="thin")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
YELLOW = PatternFill("solid", fgColor="FFFFFF00")
GREY = PatternFill("solid", fgColor="FFD9D9D9")
RED = "FFC00000"


def lines(text, width):
    n = 0
    for part in str(text).split("\n"):
        w = sum(2 if ord(ch) > 0x2E80 else 1 for ch in part)
        n += max(1, -(-w // max(int(width * 1.05), 1)))
    return n


wb = openpyxl.load_workbook(XLSX)
if SHEET in wb.sheetnames:
    del wb[SHEET]
ws = wb.create_sheet(SHEET)
ws.sheet_view.zoomScale = 70
widths = {"A": 12, "B": 52, "C": 32}
for k, v in widths.items():
    ws.column_dimensions[k].width = v

ws.merge_cells("A1:C1")
ws["A1"] = spec["title"]
ws["A1"].font = Font(name=FONT, size=14, bold=True)
ws.row_dimensions[1].height = 26
ws.merge_cells("A2:C2")
ws["A2"] = spec["intro"]
ws["A2"].font = Font(name=FONT, size=10)
ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[2].height = 34

r = 3
for c, h in enumerate(["项目", "要求", "说明"], 1):
    cell = ws.cell(r, c, h)
    cell.font = Font(name=FONT, size=10, bold=True)
    cell.fill = YELLOW
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[r].height = 20

for sec in spec["sections"]:
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws.cell(r, 1, sec["title"]).font = Font(name=FONT, size=10, bold=True)
    for c in range(1, 4):
        ws.cell(r, c).fill = GREY
        ws.cell(r, c).border = BORDER
    ws.row_dimensions[r].height = 20
    for item, req, note, key in sec["items"]:
        r += 1
        for c, v in enumerate((item, req, note), 1):
            cell = ws.cell(r, c, v or None)
            cell.font = Font(name=FONT, size=10, bold=(c == 1), color=RED if key else None)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="center",
                                       horizontal="center" if c == 1 else "left")
        n = max(lines(req, widths["B"]), lines(note, widths["C"]))
        ws.row_dimensions[r].height = max(18, n * 15 + 4)

# full drawing pages, absolutely positioned to the right of the notes table
x0 = sum(int(w * 7 + 5) for w in widths.values()) + 30
y = 10
DISPLAY_W = 1500
for i, page in enumerate(pymupdf.open(PDF)):
    longest = max(page.rect.width, page.rect.height) / 72  # inches
    png = f"{IMGDIR}\\{SHEET}_p{i + 1}.png"
    page.get_pixmap(dpi=max(72, min(200, int(4500 / longest)))).save(png)  # ~4500 px long side
    # drawings have few colours: a palette PNG is several times smaller than true colour
    PILImage.open(png).convert("RGB").quantize(colors=64, method=PILImage.Quantize.MEDIANCUT).save(png, optimize=True)
    img = XLImage(png)
    w, h = DISPLAY_W, int(img.height * DISPLAY_W / img.width)
    img.width, img.height = w, h
    img.anchor = AbsoluteAnchor(pos=XDRPoint2D(pixels_to_EMU(x0), pixels_to_EMU(y)),
                                ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h)))
    ws.add_image(img)
    y += h + 30

ws.freeze_panes = "A4"
ws.page_setup.orientation = "landscape"

# 询价单 sheet: link the drawing number, optionally rewrite material/finish/remark
main = wb.worksheets[0]
tag = f"图纸要求中文说明见 sheet「{SHEET}」"
m = spec.get("main", {})
for row in range(3, main.max_row + 1):
    if str(main.cell(row, 2).value or "").strip() != DRW_NO:
        continue
    b = main.cell(row, 2)
    # internal link: "location" (a "#Sheet!A1" target is written as an external URL, which WPS won't follow)
    b.hyperlink = Hyperlink(ref=b.coordinate, location=f"'{SHEET}'!A1", display=DRW_NO)
    b.font = Font(name="Arial", size=10.5, color="FF0563C1", underline="single")
    if main.cell(row, 5).value != "同上":
        if m.get("material"):
            main.cell(row, 5).value = m["material"]
        if m.get("finish"):
            main.cell(row, 6).value = m["finish"]
        if m.get("remark"):
            main.cell(row, 11).value = m["remark"]
    k = main.cell(row, 11)
    if tag not in (k.value or ""):
        k.value = (k.value + "\n" if k.value else "") + tag
    main.row_dimensions[row].height = max(main.row_dimensions[row].height or 30, 110 if m else 60)

wb.save(XLSX)
print("saved", XLSX, "sheet", SHEET, "rows", r)
