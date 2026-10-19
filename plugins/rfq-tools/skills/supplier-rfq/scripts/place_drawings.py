"""Final pass: put every drawing under the table and notes on the 询价单 sheet (sheet 1).

Usage: python place_drawings.py <询价单.xlsx> [more xlsx …]

Run it after all drawing sheets are added (add_drawing_sheet.py). It takes the drawing images from those sheets, in the order
their drawing numbers first appear in column B. Each drawing gets a header row ("图纸 n：<no>（询价单序号 …）", linked
to its sheet), a page break before it, and one picture per page. Each picture is fitted into a box as wide as columns
A–K and at most MAX_H px tall at 100 % zoom (about one screen), centred, and anchored "move and size with cells" over
rows sized to it, so it behaves like a picture embedded in the cells in both Excel and WPS.
Re-running rebuilds the whole block, so it is safe to run again after any change.
"""
import io
import math
import os
import sys

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Alignment, Font
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.pagebreak import Break
from PIL import Image

MARK = "图纸（中文标注版）"   # first cell of the drawing block; everything from here down is rebuilt
MAX_H = 850                # px, display height cap at 100 % zoom
SHARP = 2.0                # stored pixels per displayed pixel, so zooming in stays readable
ROW_PX = 300               # max px per picture row (Excel caps a row at 409 pt)
COLS = 11                  # A–K
FONT = "微软雅黑"


def col_px(ws, c):
    w = ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width or 8.38
    return int(w * 8 + 5)  # 等线 11 default font: 8 px per character + 5 px padding


def marker(ws, x):
    """AnchorMarker for horizontal pixel position x from the left edge of column A, row given later."""
    c = 1
    while c < COLS and x >= col_px(ws, c):
        x -= col_px(ws, c)
        c += 1
    return c - 1, pixels_to_EMU(int(x))


def shrink(data, w, h):
    """Resample to SHARP x the display size and store as a palette PNG (drawings have few colours)."""
    im = Image.open(io.BytesIO(data)).convert("RGB")
    tw = min(im.width, int(w * SHARP))
    if tw < im.width:
        im = im.resize((tw, int(im.height * tw / im.width)), Image.LANCZOS)
    im = im.quantize(colors=64, method=Image.Quantize.MEDIANCUT)
    buf = io.BytesIO()
    im.save(buf, "PNG", optimize=True)
    buf.seek(0)
    return buf


for path in sys.argv[1:]:
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]

    # find the start of the block (old block, else two rows under the last used row) and clear it
    start = next((r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == MARK), None)
    if start is None:
        last = max((c.row for row in ws.iter_rows() for c in row if c.value not in (None, "")), default=2)
        start = last + 2
    for rng in [m for m in ws.merged_cells.ranges if m.min_row >= start]:
        ws.unmerge_cells(str(rng))
    if ws.max_row >= start:
        ws.delete_rows(start, ws.max_row - start + 1)
    for r in [r for r in ws.row_dimensions if r >= start]:
        del ws.row_dimensions[r]
    ws.row_breaks.brk = [b for b in ws.row_breaks.brk if b.id < start - 1]
    ws._images = []

    # drawing sheets in the order of their links in column B, plus the line numbers that use them
    order, seqs = [], {}
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 1).value != "=ROW()-2":
            continue
        link = ws.cell(r, 2).hyperlink
        loc = (link.location or "") if link else ""
        name = loc.split("!")[0].strip("'")
        if name in wb.sheetnames:
            if name not in seqs:
                order.append(name)
            seqs.setdefault(name, []).append(r - 2)
    order += [s.title for s in wb.worksheets[1:] if s.title not in seqs and s._images]

    table_w = sum(col_px(ws, c) for c in range(1, COLS + 1))
    r = start
    ws.cell(r, 1, MARK).font = Font(name=FONT, size=12, bold=True)
    ws.cell(r, 3, "图上红框为对应的询价单序号；点击图纸号可跳到该图纸的中文说明。").font = Font(name=FONT, size=10)
    ws.row_dimensions[r].height = 22
    n_pics = 0
    for k, name in enumerate(order, 1):
        sheet = wb[name]
        r += 1
        ws.row_breaks.append(Break(id=r - 1))  # each drawing starts on a new printed page
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=COLS)
        nums = "、".join(map(str, seqs.get(name, []))) or "—"
        h = ws.cell(r, 1, f"图纸 {k}：{name}（询价单序号 {nums}）　→ 中文说明见 sheet「{name}」")
        h.hyperlink = Hyperlink(ref=h.coordinate, location=f"'{name}'!A1", display=h.value)
        h.font = Font(name=FONT, size=11, bold=True, color="FF0563C1", underline="single")
        h.alignment = Alignment(vertical="center")
        ws.row_dimensions[r].height = 24
        pics = sorted(sheet._images, key=lambda im: im.anchor._from.row if hasattr(im.anchor, "_from")
                      else im.anchor.pos.y)
        for pic in pics:
            ref = pic.ref
            data = ref.getvalue() if hasattr(ref, "getvalue") else open(ref, "rb").read()  # _data() would close the stream
            iw, ih = Image.open(io.BytesIO(data)).size
            s = min(table_w / iw, MAX_H / ih)
            w, hgt = int(iw * s), int(ih * s)
            nrows = max(1, math.ceil(hgt / ROW_PX))
            r += 1
            for i in range(nrows):
                ws.row_dimensions[r + i].height = hgt / nrows * 0.75
            img = XLImage(shrink(data, w, hgt))
            img.width, img.height = w, hgt
            x0 = (table_w - w) // 2
            c0, o0 = marker(ws, x0)
            c1, o1 = marker(ws, x0 + w)
            img.anchor = TwoCellAnchor(editAs="twoCell",
                                       _from=AnchorMarker(col=c0, colOff=o0, row=r - 1, rowOff=0),
                                       to=AnchorMarker(col=c1, colOff=o1, row=r - 1 + nrows, rowOff=0))
            ws.add_image(img)
            n_pics += 1
            r += nrows
            ws.row_dimensions[r].height = 8  # small gap under each page
    tmp = path + ".tmp"
    wb.save(tmp)  # a failed save must not break the 询价单
    os.replace(tmp, path)
    print(f"placed {len(order)} drawings ({n_pics} pages) under the notes on sheet 1 of {path}")
