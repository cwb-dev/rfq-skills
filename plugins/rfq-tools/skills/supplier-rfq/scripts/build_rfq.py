"""Build the base 询价单 sheet (one xlsx per group) from the template, before drawing sheets are added.
Usage: python build_rfq.py <groups.json> <out_dir> <rfq no>
groups.json: {"A-螺钉": {"unit": "元/千件",
                         "rows": [[B 图纸号, C 描述, D 数量, E 材质, F 表面处理, K 备注], ...],
                         "notes": ["请于 … 前回复。", ...]}, ...}
A second quantity tier of the same part: [B, "同上", qty, "同上", "同上", "同一零件第二个数量档，请分别报价"].
Output: <out_dir>/<rfq no>询价单-<group>.xlsx
"""
import copy
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

TPL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "templates", "inquiry-template.xlsx")
groups = json.load(open(sys.argv[1], encoding="utf-8"))
OUT, NO = sys.argv[2], sys.argv[3]
FONT = "微软雅黑"
thin = Side(style="thin")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WIDTHS = {"A": 6, "B": 18, "C": 30, "D": 11, "E": 26, "F": 36, "G": 12, "H": 11, "I": 10, "J": 10, "K": 42}

for name, g in groups.items():
    wb = openpyxl.load_workbook(TPL)
    ws = wb.worksheets[0]
    ws["K1"] = f"价格：人民币含税含运上海，{g['unit']}"
    ws["K1"].alignment = Alignment(wrap_text=True, vertical="center")
    tpl_a3 = copy.copy(ws["A3"].font)
    for k, v in WIDTHS.items():
        ws.column_dimensions[k].width = v
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 30
    r = 3
    for b, c, d, e, f, k in g["rows"]:
        vals = {1: "=ROW()-2", 2: b, 3: c, 4: d, 5: e, 6: f, 11: k or None}
        for col in range(1, 12):
            cell = ws.cell(r, col, vals.get(col))
            cell.border = BORDER
            cell.font = Font(name="Arial", size=10.5) if col == 2 else Font(name=FONT, size=11)
            left = col in (3, 5, 6, 11)
            cell.alignment = Alignment(wrap_text=True, vertical="center",
                                       horizontal="left" if left else "center")
        ws.cell(r, 4).number_format = "#,##0"
        r += 1
    r += 1
    ws.cell(r, 1, "说明").font = Font(name=FONT, size=10, bold=True)
    for i, note in enumerate(g["notes"], 1):
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=11)
        cell = ws.cell(r, 1, f"{i}. {note}")
        cell.font = Font(name=FONT, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 20 if len(note) < 95 else 34
    ws.freeze_panes = "A3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    path = os.path.join(OUT, f"{NO}询价单-{name}.xlsx")
    wb.save(path)
    print("saved", path, "rows", len(g["rows"]))
